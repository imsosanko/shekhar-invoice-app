#!/usr/bin/env python3
"""
Builds an .xlsx export of invoice data for Shekhar Compute Tech Services.

Invoked by the Node/Express backend (server/routes/invoices.routes.js) via
child_process.spawn: JSON describing {invoices, settings} is piped in on
stdin (invoices already filtered by search/status/date on the Node side),
and the finished .xlsx bytes are written to stdout. Nothing else must be
printed to stdout, or the file stream gets corrupted.
"""
import sys
import json
import io

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

HEADERS = [
    'Invoice No.', 'Customer', 'Customer Mobile', 'GSTIN', 'Invoice Date', 'Due Date',
    'Subtotal', 'Item Discounts', 'Additional Discount', 'CGST', 'SGST', 'IGST', 'Total Tax', 'Grand Total',
    'Status', 'Items'
]


def build_workbook(payload):
    invoices = payload.get('invoices', [])
    brand_hex = (payload.get('settings') or {}).get('primaryColor', '#1E3A8A').lstrip('#')

    wb = Workbook()
    ws = wb.active
    ws.title = 'Invoices'

    ws.append(HEADERS)
    header_fill = PatternFill(start_color=brand_hex, end_color=brand_hex, fill_type='solid')
    header_font = Font(color='FFFFFF', bold=True, size=10)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', vertical='center')

    for inv in invoices:
        customer = inv.get('customer') or {}
        items_str = '; '.join(f"{it.get('desc','')} (x{it.get('qty',0)})" for it in inv.get('items', []))
        ws.append([
            inv.get('number', ''),
            customer.get('name', ''),
            customer.get('mobile', ''),
            customer.get('gstin', ''),
            inv.get('date', ''),
            inv.get('dueDate', ''),
            round(inv.get('subtotal', 0), 2),
            round(inv.get('itemDiscountTotal', 0), 2),
            round(inv.get('additionalDiscountAmt', 0), 2),
            round(inv.get('cgst', 0), 2),
            round(inv.get('sgst', 0), 2),
            round(inv.get('igst', 0), 2),
            round(inv.get('totalTax', 0), 2),
            round(inv.get('grandTotal', 0), 2),
            inv.get('status', ''),
            items_str
        ])

    # Reasonable auto width per column based on content length.
    for idx, header in enumerate(HEADERS, start=1):
        max_len = len(header)
        for row in ws.iter_rows(min_row=2, min_col=idx, max_col=idx):
            for cell in row:
                if cell.value is not None:
                    max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[get_column_letter(idx)].width = min(45, max(12, max_len + 3))

    ws.freeze_panes = 'A2'

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def main():
    raw = sys.stdin.buffer.read()
    payload = json.loads(raw.decode('utf-8'))
    xlsx_bytes = build_workbook(payload)
    sys.stdout.buffer.write(xlsx_bytes)
    sys.stdout.buffer.flush()


if __name__ == '__main__':
    try:
        main()
    except Exception as e:
        print(f"export_excel.py error: {e}", file=sys.stderr)
        sys.exit(1)
